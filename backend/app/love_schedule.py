"""恋爱系统 · 值班日历（原生实现）。

只读最新排班表 Excel，得到「哪天上班 / 哪天休息」，并把单元格里的当天值乘车次
作为 `duty` 字段返回；前端在当天备注位置直接显示。用户备注存在
`<个人数据库>/.torch-os/love/schedule-notes.json`，可以追加交接事项或改车次。

车次时刻表、实时位置与沿线天气已经拆分到独立的工作台「列车追踪」组件，
这里不再读取 train_timetable.json。排班对象姓名只从本机配置读取
（LOVE_TARGET_NAME），不写进仓库，也不在接口里回显。
"""
from __future__ import annotations

import calendar
import json
import os
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import APIRouter, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, ConfigDict, Field, field_validator

from .vault_connection import VaultReadOnlyError, ensure_vault_writable

MAX_EXCEL_BYTES = 20 * 1024 * 1024
MAX_NOTE_LENGTH = 500
NOTES_NAME = "schedule-notes.json"
WEEKDAYS = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
TZ = ZoneInfo("Asia/Shanghai")


def _find_latest_excel(base_dir: Path) -> Path | None:
    base_dir = Path(base_dir)
    if not base_dir.is_dir():
        return None
    files = [path for path in base_dir.glob("*.xlsx") if path.is_file()]
    files += [path for path in base_dir.glob("*.xlsm") if path.is_file()]
    if not files:
        return None
    files.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    for path in files:
        if any(hint in path.name for hint in ("排班", "值班", "交路")):
            return path
    return files[0]


def _is_rest_cell(cell_value: Any) -> bool:
    if cell_value is None:
        return True
    text = str(cell_value).strip()
    return not text or text in ("休", "公休") or text.startswith("段抽")


def _duty_text(cell_value: Any) -> str:
    """值班表单元格里的当天值乘车次；休/公休/请假等没有车次时返回空。"""
    if cell_value is None or _is_rest_cell(cell_value):
        return ""
    text = re.sub(r"\s+", " ", str(cell_value)).strip()
    # 只把真正含车次号的单元格当成“当天值乘”，避免把育儿假、疗养等当车次。
    if not re.search(r"[A-Za-z]\d{1,5}", text):
        return ""
    return text[:MAX_NOTE_LENGTH]


def _excel_serial_date(serial: int) -> date:
    return date(1899, 12, 30) + timedelta(days=int(serial))


def _looks_like_schedule(sheet: Any) -> bool:
    if sheet is None:
        return False
    found = 0
    for column in range(10, 21):
        day_number = sheet.cell(row=7, column=column).value
        serial = sheet.cell(row=8, column=column).value
        try:
            int(day_number)
            int(serial)
        except (TypeError, ValueError):
            continue
        found += 1
        if found >= 2:
            return True
    return False


class ScheduleWorkbook:
    """排班表读取：日期列按「日号递推」而不是直接相信序列号（沿用原项目踩坑后的算法）。"""

    def __init__(self, path: Path | None, target_name: str = ""):
        self.path = Path(path) if path else None
        self.target_name = str(target_name or "").strip()
        self.workbook = None
        self.sheet = None
        self.date_columns: dict[date, str] = {}
        if self.path:
            self._load()

    def _load(self) -> None:
        import openpyxl

        self.workbook = openpyxl.load_workbook(str(self.path), data_only=True)
        self.sheet = self._pick_sheet()
        self._build_dates()

    def _pick_sheet(self):
        if self.workbook is None:
            return None
        schedules = [sheet for sheet in self.workbook.worksheets if _looks_like_schedule(sheet)]
        if not schedules:
            return None
        if not self.target_name:
            return schedules[0]
        for sheet in schedules:
            for row in range(7, min(sheet.max_row, 300) + 1):
                value = sheet.cell(row=row, column=5).value
                if value is not None and str(value).strip() == self.target_name:
                    return sheet
        return None

    def _build_dates(self) -> None:
        import openpyxl

        if self.sheet is None:
            return
        rows: list[tuple[int, int, int]] = []
        for column in range(10, 61):
            day_number = self.sheet.cell(row=7, column=column).value
            serial = self.sheet.cell(row=8, column=column).value
            if day_number is None or serial is None:
                continue
            try:
                rows.append((column, int(day_number), int(serial)))
            except (TypeError, ValueError):
                continue
        if not rows:
            return
        anchor = next(((col, day, serial) for col, day, serial in rows if day == 1), None)
        if anchor is None:
            for col, _day, serial in rows:
                try:
                    self.date_columns[_excel_serial_date(serial)] = openpyxl.utils.get_column_letter(col)
                except (TypeError, ValueError):
                    continue
            return
        anchor_col, _anchor_day, anchor_serial = anchor
        try:
            anchor_date = _excel_serial_date(anchor_serial)
        except (TypeError, ValueError):
            return
        first_day = rows[0][1]
        if first_day == 1:
            # 表从某月 1 号开始：第一行就是当月起点，不能减到上个月。
            current_month_start = date(anchor_date.year, anchor_date.month, 1)
        else:
            previous_month = anchor_date.month - 1
            previous_year = anchor_date.year
            if previous_month == 0:
                previous_month = 12
                previous_year -= 1
            days_in_previous = calendar.monthrange(previous_year, previous_month)[1]
            current_month_start = anchor_date - timedelta(days=days_in_previous)
        previous_day = 0
        for col, day_number, _serial in rows:
            if previous_day > 0 and day_number < previous_day:
                days_in_month = calendar.monthrange(
                    current_month_start.year, current_month_start.month
                )[1]
                current_month_start = current_month_start + timedelta(days=days_in_month)
            previous_day = day_number
            self.date_columns[current_month_start + timedelta(days=day_number - 1)] = (
                openpyxl.utils.get_column_letter(col)
            )

    def person_rows(self) -> list[tuple[int, str]]:
        """排班表 E 列里的姓名行；只在本机内存里用于定位，不回给前端。"""
        if self.sheet is None:
            return []
        rows: list[tuple[int, str]] = []
        for row in range(7, self.sheet.max_row + 1):
            value = self.sheet.cell(row=row, column=5).value
            name = str(value).strip() if value is not None else ""
            if name:
                rows.append((row, name))
        return rows

    def find_person_row(self) -> int | None:
        rows = self.person_rows()
        if not rows:
            return None
        if self.target_name:
            for row, name in rows:
                if name == self.target_name:
                    return row
            return None
        return rows[0][0] if len(rows) == 1 else None

    def cell_for(self, row: int, column_letter: str) -> Any:
        import openpyxl

        if self.sheet is None:
            return None
        return self.sheet.cell(
            row=row, column=openpyxl.utils.column_index_from_string(column_letter)).value

    def date_range(self) -> tuple[date, date] | None:
        if not self.date_columns:
            return None
        dates = sorted(self.date_columns)
        return dates[0], dates[-1]


# ── 备注 ─────────────────────────────────────────────────────────────

def _notes_path(notes_path: Path | str) -> Path:
    return Path(notes_path)


def _load_notes(path: Path) -> dict[str, str]:
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return {}
    if not isinstance(data, dict):
        return {}
    notes: dict[str, str] = {}
    for key, value in data.items():
        if not isinstance(key, str) or not isinstance(value, str):
            continue
        try:
            date.fromisoformat(key)
        except ValueError:
            continue
        if value.strip():
            notes[key] = value.strip()[:MAX_NOTE_LENGTH]
    return notes


def _write_notes(path: Path, notes: dict[str, str], *, read_only: bool) -> None:
    try:
        ensure_vault_writable(read_only, "保存值班备注")
    except VaultReadOnlyError as error:
        raise HTTPException(409, str(error)) from error
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(notes, ensure_ascii=False, indent=2), encoding="utf-8", newline="\n")
    os.replace(temporary, path)


def _month_days(month: str) -> list[date]:
    year, month_number = (int(part) for part in month.split("-"))
    days = calendar.monthrange(year, month_number)[1]
    return [date(year, month_number, day) for day in range(1, days + 1)]


def _month_payload(month: str, excel_dir: Path, notes: dict[str, str],
                   target_name: str) -> dict[str, Any]:
    month_days = _month_days(month)
    base = {"month": month, "today": datetime.now(TZ).date().isoformat()}
    latest = _find_latest_excel(excel_dir)
    if latest is None:
        return {**base, "found": False, "error": "还没有排班表", "days": []}
    try:
        workbook = ScheduleWorkbook(latest, target_name)
    except Exception as exc:  # noqa: BLE001 - 第三方 Excel 解析错误要转成界面提示
        return {**base, "found": False, "error": f"排班表读取失败：{exc}", "days": []}
    if workbook.sheet is None:
        return {
            **base, "found": False,
            "error": "排班表里没有找到可识别的日期列或排班对象，请检查文件与 LOVE_TARGET_NAME。",
            "days": [],
        }
    row = workbook.find_person_row()
    if row is None:
        error = ("排班表里有多个姓名，请在本机 .env 设置 LOVE_TARGET_NAME 后重启服务。"
                 if len(workbook.person_rows()) > 1 else "排班表里没有找到配置的排班对象，请检查 LOVE_TARGET_NAME。")
        return {**base, "found": False, "error": error, "days": []}
    days = []
    for day in month_days:
        column = workbook.date_columns.get(day)
        note = notes.get(day.isoformat(), "")
        if column is None:
            days.append({"date": day.isoformat(), "in_range": False,
                         "work": None, "rest": None, "note": note, "duty": ""})
            continue
        cell = workbook.cell_for(row, column)
        rest = _is_rest_cell(cell)
        days.append({"date": day.isoformat(), "in_range": True,
                     "work": not rest, "rest": rest, "note": note,
                     "duty": "" if rest else _duty_text(cell)})
    span = workbook.date_range()
    return {
        **base,
        "found": True,
        "days": days,
        "range": (
            {"start": span[0].isoformat(), "end": span[1].isoformat()}
            if span else None
        ),
    }


# ── API ──────────────────────────────────────────────────────────────

class NoteWrite(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    date: str = Field(min_length=10, max_length=10)
    note: str = Field(default="", max_length=MAX_NOTE_LENGTH)

    @field_validator("date")
    @classmethod
    def validate_date(cls, value: str) -> str:
        try:
            date.fromisoformat(value)
        except ValueError as error:
            raise ValueError("日期格式应为 YYYY-MM-DD") from error
        return value


def create_schedule_router(excel_dir: Path, notes_path: Path,
                           target_name: str = "", *, read_only: bool = False,
                           prefix: str = "/api/love/schedule") -> APIRouter:
    excel_dir = Path(excel_dir)
    notes_path = _notes_path(notes_path)
    target_name = str(target_name or "").strip()
    router = APIRouter(prefix=prefix, tags=["love-schedule"])

    @router.get("/status")
    def status():
        latest = _find_latest_excel(excel_dir)
        span = None
        if latest is not None:
            try:
                workbook = ScheduleWorkbook(latest, target_name)
                dates = workbook.date_range()
                if dates:
                    span = {"start": dates[0].isoformat(), "end": dates[1].isoformat()}
            except Exception:  # noqa: BLE001
                span = None
        return {
            "excel_dir": str(excel_dir),
            "notes_path": str(notes_path),
            "excel_exists": latest is not None,
            "target_configured": bool(target_name),
            "date_range": span,
        }

    @router.get("/month")
    def month(month: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")):
        return _month_payload(month, excel_dir, _load_notes(notes_path), target_name)

    @router.get("/day")
    def day(target_date: str = Query(alias="date", pattern=r"^\d{4}-\d{2}-\d{2}$")):
        try:
            parsed = date.fromisoformat(target_date)
        except ValueError as error:
            raise HTTPException(422, "日期格式应为 YYYY-MM-DD") from error
        payload = _month_payload(parsed.strftime("%Y-%m"), excel_dir,
                                 _load_notes(notes_path), target_name)
        if not payload.get("found"):
            return {"found": False, "date": target_date, "error": payload.get("error", "排班表不可用")}
        item = next((day for day in payload["days"] if day["date"] == target_date), None)
        return {"found": True, "date": target_date, "weekday": WEEKDAYS[parsed.weekday()],
                **(item or {"in_range": False, "work": None, "rest": None,
                            "note": "", "duty": ""})}

    @router.get("/upcoming")
    def upcoming(days: int = Query(default=7, ge=1, le=31),
                 start: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}-\d{2}$")):
        """工作台值班小组件：从今天（或 start）起连续几天。"""
        if start:
            try:
                first = date.fromisoformat(start)
            except ValueError as error:
                raise HTTPException(422, "日期格式应为 YYYY-MM-DD") from error
        else:
            first = datetime.now(TZ).date()
        dates = [first + timedelta(days=offset) for offset in range(days)]
        grouped: dict[str, list[date]] = {}
        for item in dates:
            grouped.setdefault(item.strftime("%Y-%m"), []).append(item)
        items: list[dict[str, Any]] = []
        found = False
        error = ""
        notes = _load_notes(notes_path)
        for month, month_dates in grouped.items():
            payload = _month_payload(month, excel_dir, notes, target_name)
            if not payload.get("found"):
                error = error or str(payload.get("error") or "排班表不可用")
                for item in month_dates:
                    items.append({
                        "date": item.isoformat(), "weekday": WEEKDAYS[item.weekday()],
                        "in_range": False, "work": None, "rest": None,
                        "note": notes.get(item.isoformat(), ""), "duty": "",
                    })
                continue
            found = True
            by_date = {entry["date"]: entry for entry in payload.get("days", [])}
            for item in month_dates:
                entry = by_date.get(item.isoformat()) or {
                    "date": item.isoformat(), "in_range": False, "work": None,
                    "rest": None, "note": notes.get(item.isoformat(), ""), "duty": "",
                }
                items.append({**entry, "weekday": WEEKDAYS[item.weekday()]})
        return {
            "today": datetime.now(TZ).date().isoformat(),
            "start": first.isoformat(),
            "found": found,
            "error": error,
            "days": items,
        }

    @router.put("/note")
    def save_note(payload: NoteWrite):
        notes = _load_notes(notes_path)
        note = payload.note.strip()
        if note:
            notes[payload.date] = note
        else:
            notes.pop(payload.date, None)
        _write_notes(notes_path, notes, read_only=read_only)
        return {"date": payload.date, "note": note}

    @router.post("/excel")
    async def replace_excel(file: UploadFile = File(...)):
        import openpyxl

        try:
            ensure_vault_writable(read_only, "导入排班表")
        except VaultReadOnlyError as error:
            raise HTTPException(409, str(error)) from error
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in (".xlsx", ".xlsm"):
            raise HTTPException(422, "仅支持 .xlsx / .xlsm 格式")
        raw = await file.read(MAX_EXCEL_BYTES + 1)
        if len(raw) > MAX_EXCEL_BYTES:
            raise HTTPException(413, "排班表超过 20MB 上限")
        if not raw:
            raise HTTPException(400, "文件内容为空")
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
                handle.write(raw)
                temporary_path = Path(handle.name)
            try:
                workbook = openpyxl.load_workbook(str(temporary_path), data_only=True)
            except Exception as exc:  # noqa: BLE001
                raise HTTPException(422, f"无法读取 Excel：{exc}") from exc
            try:
                valid = any(_looks_like_schedule(sheet) for sheet in workbook.worksheets)
            finally:
                workbook.close()
            if not valid:
                raise HTTPException(422, "排班表里找不到日期列（第 7/8 行的日号与日期序列）")
            safe_name = re.sub(r"[^\w一-鿿.\-()（）]+", "_", Path(file.filename or "排班表.xlsx").name)[-120:]
            if not safe_name.lower().endswith(suffix):
                safe_name += suffix
            excel_dir.mkdir(parents=True, exist_ok=True)
            destination = excel_dir / safe_name
            shutil.copyfile(temporary_path, destination)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        return {"success": True, "message": "排班表已更新。"}

    return router
