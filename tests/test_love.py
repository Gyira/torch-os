"""原生恋爱系统：数据库复用、导入、统计、画像、日志、AI 编排与值班日历。"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.app import love as love_module
from backend.app import love_schedule
from backend.app.love import ChatParser, LoveStore, create_love_router

SAMPLE = """2026-05-20 10:00:00 我: 你好，我的电话是13800138000
2026-05-20 10:01:00 小林: 嗯嗯，今天做什么
2026-05-20 10:02:00 我: 明天见面吧
2026-05-21 09:00:00 小林: 好呀
2026-05-21 09:01:00 我: 那说定了
"""

AI_RESULT = {
    "relationship_stage": "稳定期",
    "emotional_tone": {"overall": "温和", "fluctuation": 0.2, "trend": "稳定"},
    "communication_patterns": {
        "initiative_balance": "双方接近",
        "response_style": "及时回应",
        "key_topics": ["见面", "日常"],
    },
    "key_events": [],
    "risk_warnings": [],
    "personalized_advice": ["保持直接沟通"],
    "partner_traits": [{
        "trait": "沟通风格", "value": "直接、回应快",
        "confidence": 0.9, "evidence": "每次都能很快接话",
    }],
}


def build_client(tmp_path: Path) -> tuple[TestClient, LoveStore]:
    store = LoveStore(tmp_path / "love.db", key_file=tmp_path / ".api_key.key")
    store.ensure_schema()
    app = FastAPI()
    app.include_router(create_love_router(store))
    return TestClient(app), store


def import_sample(client: TestClient, name: str = "测试项目") -> str:
    response = client.post("/api/love/import/text", json={
        "content": SAMPLE, "file_name": "sample.txt", "project_name": name,
    })
    assert response.status_code == 200, response.text
    return response.json()["project_id"]


def test_parser_handles_gbk_and_multiline() -> None:
    parser = ChatParser()
    raw = (
        "2026-05-20 22:19:08 '我'\n"
        "你已添加了🦋，以上是打招呼的消息。\n"
        "2026-05-20 22:20:00 '小林'\n"
        "嗨\n"
    ).encode("gb18030")
    parsed = parser.decode_and_parse(raw)
    assert len(parsed.messages) == 2
    assert parsed.senders == ["我", "小林"]
    assert parsed.messages[0].is_system is False or parsed.messages[0].sender == "我"
    # 单行与多行混排也要能解析。
    mixed = parser.parse(
        "2026-05-21 10:00:00 我: 早上好\n"
        "2026-05-21 10:01:00 '小林'\n"
        "早呀\n"
    )
    assert [message.content for message in mixed.messages] == ["早上好", "早呀"]


def test_import_stats_journal_profile_and_report(tmp_path: Path) -> None:
    client, store = build_client(tmp_path)
    assert client.get("/api/love/status").json()["available"] is True

    # 前端使用 multipart 上传；先预览再导入。
    preview = client.post(
        "/api/love/import/preview",
        files={"file": ("sample.txt", SAMPLE.encode("utf-8"), "text/plain")},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["message_count"] == 5
    assert preview.json()["senders"] == ["小林", "我"] or set(preview.json()["senders"]) == {"小林", "我"}
    uploaded = client.post(
        "/api/love/import",
        files={"file": ("sample.txt", SAMPLE.encode("gb18030"), "text/plain")},
        data={"project_name": "上传项目"},
    )
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()["message_count"] == 5
    assert uploaded.json()["project_name"].startswith("上传项目")
    project_id = uploaded.json()["project_id"]

    projects = client.get("/api/love/projects").json()["projects"]
    assert [project["id"] for project in projects] == [project_id]
    assert projects[0]["message_count"] == 5

    overview = client.get(f"/api/love/projects/{project_id}/stats/overview").json()
    assert overview["total_messages"] == 5
    assert overview["sender_a"]["messages"] == 3
    assert overview["sender_b"]["messages"] == 2
    assert overview["sender_a"]["ratio"] + overview["sender_b"]["ratio"] == pytest.approx(1)

    trends = client.get(
        f"/api/love/projects/{project_id}/stats/trends?granularity=day"
    ).json()
    assert [row["date"] for row in trends["data"]] == ["2026-05-20", "2026-05-21"]

    timing = client.get(f"/api/love/projects/{project_id}/stats/timing").json()
    assert timing["median_reply_interval_seconds"] > 0
    assert sum(timing["reply_distribution"].values()) > 0

    media = client.get(f"/api/love/projects/{project_id}/stats/media").json()
    assert media["text_count"] == 5
    keywords = client.get(f"/api/love/projects/{project_id}/stats/keywords?top_n=20").json()
    assert isinstance(keywords["words"], list)

    # 日志：新增、修改、导出。
    created = client.post(f"/api/love/projects/{project_id}/journal", json={
        "date": "2026-05-20", "content": "今天聊得很开心", "mood": "happy",
    })
    assert created.status_code == 200, created.text
    entry_id = created.json()["id"]
    updated = client.put(
        f"/api/love/projects/{project_id}/journal/{entry_id}",
        json={"content": "今天聊得很开心，还约了见面"},
    )
    assert "约了见面" in updated.json()["content"]
    export = client.get(f"/api/love/projects/{project_id}/journal/export")
    assert export.status_code == 200
    assert export.content.startswith(b"\xef\xbb\xbf")
    assert "filename*=UTF-8''" in export.headers["content-disposition"]

    # 画像：手动新增、锁定、修改、删除。
    profile = client.post(f"/api/love/projects/{project_id}/profile/traits", json={
        "trait": "雷区", "value": "迟到", "evidence": "用户观察",
    })
    assert profile.status_code == 200, profile.text
    trait = next(item for item in profile.json()["traits"] if item["trait"] == "雷区")
    assert trait["locked"] is False
    locked = client.put(
        f"/api/love/projects/{project_id}/profile/traits/{'雷区'}/lock",
        json={"locked": True},
    )
    assert locked.json()["traits"][0]["locked"] is True
    changed = client.put(
        f"/api/love/projects/{project_id}/profile/traits/{'雷区'}",
        json={"value": "迟到与失约"},
    )
    assert changed.json()["traits"][0]["value"] == "迟到与失约"
    removed = client.delete(f"/api/love/projects/{project_id}/profile/traits/{'雷区'}")
    assert all(item["trait"] != "雷区" for item in removed.json()["traits"])

    # 报告：HTML 预览走单独文档；项目名必须转义。
    client.put(f"/api/love/projects/{project_id}", json={"name": "<b>测试</b>"})
    report = client.post(f"/api/love/projects/{project_id}/report")
    assert report.status_code == 200, report.text
    html = report.json()["html"]
    assert "&lt;b&gt;测试&lt;/b&gt;" in html
    assert "<b>测试</b>" not in html
    preview = client.get(f"/api/love/projects/{project_id}/report/preview")
    assert preview.status_code == 200
    assert preview.headers["content-type"].startswith("text/html")

    # 删除项目必须显式确认，并级联删除消息和日志。
    assert client.request(
        "DELETE", f"/api/love/projects/{project_id}", json={"confirm": False}
    ).status_code == 422
    deleted = client.request(
        "DELETE", f"/api/love/projects/{project_id}", json={"confirm": True}
    )
    assert deleted.status_code == 200
    with store.connect() as con:
        assert con.execute("SELECT COUNT(*) FROM messages").fetchone()[0] == 0
        assert con.execute("SELECT COUNT(*) FROM journal_entries").fetchone()[0] == 0


def test_api_key_is_encrypted_and_analysis_runs_offline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    client, store = build_client(tmp_path)
    project_id = import_sample(client)

    # 没有 Key 时不允许启动分析，避免先把日志写成 running 再失败。
    assert client.post(f"/api/love/projects/{project_id}/analysis").status_code == 400

    plaintext = "sk-test-love-key-123456"
    saved = client.post("/api/love/settings/api-key", json={"key": plaintext})
    assert saved.status_code == 200, saved.text
    assert store.key_file.is_file() and len(store.key_file.read_bytes()) == 32
    with store.connect() as con:
        stored = con.execute("SELECT value FROM settings WHERE key='api_key'").fetchone()[0]
    assert plaintext not in stored
    assert client.get("/api/love/settings").json()["api_key_configured"] is True

    seen: list[list[dict[str, str]]] = []

    def fake_deepseek(settings, api_key: str, messages: list[dict[str, str]], **kwargs):
        assert api_key == plaintext
        seen.append(messages)
        return json.dumps(AI_RESULT, ensure_ascii=False), {
            "prompt_tokens": 12, "completion_tokens": 34,
        }

    monkeypatch.setattr(love_module, "_call_deepseek", fake_deepseek)
    started = client.post(f"/api/love/projects/{project_id}/analysis", json={"analysis_type": "deep"})
    assert started.status_code == 200, started.text
    log_id = started.json()["log_id"]

    status = client.get(f"/api/love/projects/{project_id}/analysis/status/{log_id}").json()
    assert status["status"] == "completed"
    assert status["result"]["relationship_stage"] == "稳定期"
    assert status["prompt_tokens"] == 12

    prompt = seen[0][1]["content"]
    assert "13800138000" not in prompt
    assert "小林" not in prompt
    assert "用户B" in prompt
    profile = client.get(f"/api/love/projects/{project_id}/profile").json()
    assert any(item["trait"] == "沟通风格" for item in profile["traits"])

    deleted = client.delete("/api/love/settings/api-key")
    assert deleted.status_code == 200
    assert client.get("/api/love/settings").json()["api_key_configured"] is False


def _excel_serial(day: date) -> int:
    return (day - date(1899, 12, 30)).days


def _make_schedule_fixture(tmp_path: Path) -> Path:
    import openpyxl

    excel_dir = tmp_path / "xlsx"
    excel_dir.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "高1-14组"
    sheet.cell(row=3, column=5, value="2026年5月份排班表")
    days = [date(2026, 5, 20), date(2026, 5, 21), date(2026, 5, 22),
            date(2026, 6, 20), date(2026, 6, 21), date(2026, 6, 22)]
    cells = ["G1", "休", "G2", "G3", "G4", "G5"]
    for index, (day, value) in enumerate(zip(days, cells)):
        column = 10 + index
        sheet.cell(row=7, column=column, value=day.day)
        sheet.cell(row=8, column=column, value=_excel_serial(day))
        sheet.cell(row=10, column=column, value=value)
    sheet.cell(row=9, column=3, value="高1组")
    sheet.cell(row=10, column=5, value="小林")
    excel_path = excel_dir / "5月份排班.xlsx"
    workbook.save(excel_path)
    return excel_dir


def _schedule_client(excel_dir: Path, notes_path: Path, target_name: str = "小林"):
    app = FastAPI()
    app.include_router(love_schedule.create_schedule_router(
        excel_dir, notes_path, target_name=target_name))
    return TestClient(app)


def test_schedule_duty_text_only_returns_train_cells() -> None:
    assert love_schedule._duty_text("G336/G1571") == "G336/G1571"
    assert love_schedule._duty_text("G334-C2735-C2736-G337") == "G334-C2735-C2736-G337"
    for value in ("休", "公休", "段抽", "育儿假", "疗养", "热备", "", None):
        assert love_schedule._duty_text(value) == ""


def test_schedule_month_calendar_and_notes(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    notes_path = tmp_path / "schedule-notes.json"
    client = _schedule_client(excel_dir, notes_path)

    payload = client.get("/api/love/schedule/month?month=2026-05").json()
    assert payload["found"] is True
    days = {item["date"]: item for item in payload["days"]}
    assert days["2026-05-20"]["in_range"] is True
    assert days["2026-05-20"]["work"] is True and days["2026-05-20"]["rest"] is False
    assert days["2026-05-20"]["duty"] == "G1"
    assert days["2026-05-21"]["work"] is False and days["2026-05-21"]["rest"] is True
    assert days["2026-05-21"]["duty"] == ""
    assert days["2026-05-22"]["duty"] == "G2"
    assert days["2026-05-19"]["in_range"] is False
    assert days["2026-05-19"]["duty"] == ""
    assert "trains" not in payload and "weather" not in payload
    assert payload["range"] == {"start": "2026-05-20", "end": "2026-06-22"}

    # 备注是可写的自由文本，写车次也行；空备注删除。
    saved = client.put("/api/love/schedule/note", json={"date": "2026-05-20", "note": "G1 / G2"})
    assert saved.status_code == 200, saved.text
    refreshed = client.get("/api/love/schedule/month?month=2026-05").json()
    assert {item["date"]: item for item in refreshed["days"]}["2026-05-20"]["note"] == "G1 / G2"
    cleared = client.put("/api/love/schedule/note", json={"date": "2026-05-20", "note": ""})
    assert cleared.status_code == 200
    assert notes_path.read_text(encoding="utf-8").strip() in ("{}", "")

    status = client.get("/api/love/schedule/status").json()
    assert status["target_configured"] is True
    assert "target_name" not in status
    assert "timetable" not in status
    assert status["date_range"] == {"start": "2026-05-20", "end": "2026-06-22"}


def test_schedule_day_endpoint_reports_work_and_rest(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    client = _schedule_client(excel_dir, tmp_path / "notes.json")

    work = client.get("/api/love/schedule/day?date=2026-05-20").json()
    assert work["found"] is True and work["work"] is True and work["weekday"] == "周三"
    assert work["duty"] == "G1"
    rest = client.get("/api/love/schedule/day?date=2026-05-21").json()
    assert rest["found"] is True and rest["rest"] is True and rest["duty"] == ""
    out = client.get("/api/love/schedule/day?date=2026-05-19").json()
    assert out["found"] is True and out["in_range"] is False and out["duty"] == ""


def test_schedule_upcoming_days_for_workbench_widget(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    client = _schedule_client(excel_dir, tmp_path / "notes.json")
    payload = client.get(
        "/api/love/schedule/upcoming?start=2026-05-20&days=3"
    ).json()
    assert payload["found"] is True
    assert payload["start"] == "2026-05-20"
    assert [item["date"] for item in payload["days"]] == [
        "2026-05-20", "2026-05-21", "2026-05-22",
    ]
    assert [item["duty"] for item in payload["days"]] == ["G1", "", "G2"]
    assert payload["days"][0]["weekday"] == "周三"
    assert payload["days"][1]["rest"] is True
    assert "target_name" not in payload

    outside = client.get(
        "/api/love/schedule/upcoming?start=2026-05-18&days=1"
    ).json()
    assert outside["found"] is True
    assert outside["days"][0]["in_range"] is False
    assert outside["days"][0]["duty"] == ""


def test_schedule_auto_detects_single_person_without_configured_name(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    client = _schedule_client(excel_dir, tmp_path / "notes.json", target_name="")
    payload = client.get("/api/love/schedule/month?month=2026-05").json()
    assert payload["found"] is True
    days = {item["date"]: item for item in payload["days"]}
    assert days["2026-05-20"]["work"] is True


def test_schedule_asks_for_name_instead_of_guessing_between_people(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    import openpyxl

    path = excel_dir / "5月份排班.xlsx"
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=11, column=5, value="另一个姓名")
    workbook.save(path)
    client = _schedule_client(excel_dir, tmp_path / "notes.json", target_name="")

    payload = client.get("/api/love/schedule/month?month=2026-05").json()
    assert payload["found"] is False
    assert "LOVE_TARGET_NAME" in payload["error"]
    # 错误信息只提示配置，不把排班表里的真实姓名回给界面。
    assert "小林" not in payload["error"]
    assert "另一个姓名" not in payload["error"]


def test_schedule_excel_upload_validates_the_workbook(tmp_path: Path) -> None:
    excel_dir = _make_schedule_fixture(tmp_path)
    client = _schedule_client(excel_dir, tmp_path / "notes.json")
    payload = (excel_dir / "5月份排班.xlsx").read_bytes()
    response = client.post(
        "/api/love/schedule/excel",
        files={"file": ("新排班.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["message"] == "排班表已更新。"

    import openpyxl
    bad_path = tmp_path / "bad.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.cell(row=1, column=1, value="not a schedule")
    workbook.save(bad_path)
    bad = client.post(
        "/api/love/schedule/excel",
        files={"file": ("bad.xlsx", bad_path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert bad.status_code == 422
